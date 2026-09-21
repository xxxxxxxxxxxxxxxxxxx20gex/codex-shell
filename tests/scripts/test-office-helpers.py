"""Run with a Python environment containing CS Office requirements."""

import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.dont_write_bytecode = True

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "bundled/office-marketplace/plugins/cs-office/scripts"


def load(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


render_office = load("render_office")
templates = load("create_template")


class OfficeHelpersTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="cs-office-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        import pymupdf
        self.pdf = self.root / "中文 参考.pdf"
        with pymupdf.open() as document:
            page = document.new_page(width=300, height=200)
            page.insert_text((30, 50), "CS Office reference")
            document.save(self.pdf)
        self.source_bytes = self.pdf.read_bytes()

    def test_pdf_renders_and_does_not_modify_source(self):
        result = render_office.render(self.pdf, self.root / "预览")
        self.assertEqual(len(result["pages"]), 1)
        import pymupdf
        pixmap = pymupdf.Pixmap(result["pages"][0])
        self.assertEqual((pixmap.width, pixmap.height), (600, 400))
        self.assertGreater(len(set(pixmap.samples)), 1)
        self.assertEqual(self.pdf.read_bytes(), self.source_bytes)
        with self.assertRaises(ValueError):
            render_office.render(self.pdf, self.root / "预览")

    def test_converter_failure_does_not_report_success(self):
        from docx import Document
        source = self.root / "报告.docx"
        Document().save(source)
        with patch.object(render_office.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, b"", b"")):
            with self.assertRaisesRegex(RuntimeError, "转换失败"):
                render_office.render(source, self.root / "failed", sys.executable)
        with patch.object(render_office.shutil, "which", return_value=None):
            with self.assertRaises(FileNotFoundError):
                render_office.render(source, self.root / "missing")

    def test_private_converter_profile_and_page_output(self):
        from pptx import Presentation
        source = self.root / "演示.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0])
        deck.save(source)
        output = self.root / "slides"

        def convert(command, **kwargs):
            self.assertIn("--headless", command)
            self.assertIn("-env:UserInstallation=file:", command[1])
            self.assertEqual(kwargs["timeout"], 120)
            (output / "演示.pdf").write_bytes(self.source_bytes)
            return subprocess.CompletedProcess(command, 0, b"", b"")

        with patch.object(render_office.subprocess, "run", side_effect=convert):
            result = render_office.render(source, output, sys.executable)
        self.assertTrue(Path(result["pages"][0]).is_file())
        self.assertEqual(len(Presentation(source).slides), 1)

    def test_template_preserves_reference_and_refuses_overwrite(self):
        import yaml
        preview = Path(render_office.render(self.pdf, self.root / "preview")["pages"][0])
        with patch.dict(os.environ, {"CODEX_HOME": str(self.root / "cs-home")}):
            target = templates.create_template(self.pdf, preview, "report", "报告模板", "创建报告：保留参考版式", "保留标题层级。")
            self.assertEqual((target / "assets/reference.pdf").read_bytes(), self.source_bytes)
            text = (target / "SKILL.md").read_text(encoding="utf-8")
            self.assertIn("保留标题层级。", text)
            metadata = yaml.safe_load((target / "agents/openai.yaml").read_text(encoding="utf-8"))
            self.assertEqual(metadata["interface"]["display_name"], "报告模板")
            with self.assertRaises(FileExistsError):
                templates.create_template(self.pdf, preview, "report", "报告", "用途")
            self.assertFalse(list(target.parent.glob(".cs-template-*")))

    def test_template_rejects_unsafe_names_missing_home_and_invalid_preview(self):
        preview = Path(render_office.render(self.pdf, self.root / "preview")["pages"][0])
        with patch.dict(os.environ, {"CODEX_HOME": ""}):
            with self.assertRaises(ValueError):
                templates.create_template(self.pdf, preview, "report", "报告", "用途")
        with patch.dict(os.environ, {"CODEX_HOME": str(Path.home() / ".codex")}):
            with self.assertRaises(ValueError):
                templates.create_template(self.pdf, preview, "report", "报告", "用途")
        with self.assertRaises(ValueError):
            templates.create_template(self.pdf, preview, "../escape", "报告", "用途")
        with self.assertRaises(ValueError):
            templates.create_template(self.pdf, self.pdf, "report", "报告", "用途")

    def test_office_templates_keep_original_packages(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook, load_workbook
        from zipfile import ZipFile
        preview = Path(render_office.render(self.pdf, self.root / "preview")["pages"][0])
        document = Document()
        document.add_heading("报告", level=1)
        document.save(self.root / "source.docx")
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0])
        deck.save(self.root / "source.pptx")
        workbook = Workbook()
        workbook.active["A1"] = "=1+2"
        workbook.save(self.root / "source.xlsx")
        with patch.dict(os.environ, {"CODEX_HOME": str(self.root / "cs-home")}):
            for extension in ("docx", "pptx", "xlsx"):
                source = self.root / f"source.{extension}"
                target = templates.create_template(source, preview, extension, "模板", "按原版式创建文件")
                retained = target / f"assets/reference.{extension}"
                self.assertEqual(retained.read_bytes(), source.read_bytes())
                with ZipFile(retained) as package:
                    self.assertIsNone(package.testzip())
            retained = self.root / "cs-home/skills/cs-template-xlsx/assets/reference.xlsx"
            self.assertEqual(load_workbook(retained).active["A1"].value, "=1+2")

    def test_all_dependency_modes_report_structured_results(self):
        for mode in ("pdf", "documents", "spreadsheets", "presentations", "templates"):
            result = subprocess.run([sys.executable, str(SCRIPTS / "check_dependencies.py"), mode], capture_output=True, check=False)
            data = json.loads(result.stdout)
            self.assertEqual(data["mode"], mode)
            self.assertIn("pymupdf", data["renderModules"])
            self.assertEqual(result.returncode, 0 if all(data["requiredModules"].values()) else 1)


if __name__ == "__main__":
    unittest.main()
